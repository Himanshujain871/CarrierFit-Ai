import io
import docx

def extract_text_from_docx_bytes(docx_bytes: bytes) -> str:
    """Extract raw text from DOCX file bytes."""
    try:
        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
        
        # Also extract table text if present
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))
                    
        return "\n".join(full_text)
    except Exception as e:
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}")
